# -*- coding: utf-8 -*-
"""Planilha para o Ricardo separar cada variedade por criadouro."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

rows = json.load(open('/tmp/variedades.json', encoding='utf-8'))
CAT = {'aquaticas':'Aquáticas','galiformes':'Galiformes','pombas':'Pombas','psitacideos':'Psitacídeos',
       'turacos':'Turacos','exoticas':'Exóticas','faisoes':'Faisões','pavoes':'Pavões','perdizes':'Perdizes',
       'ring-necks':'Ring Necks','loris':'Lóris','francolins':'Francolins','sandgrouse':'Sandgrouse'}
ordem = {'aves-arca': 0, 'parceiros': 1}
rows.sort(key=lambda r: (ordem.get(r['criador'], 2), r['grupo'], r['nome']))

VERDE = '1F3B2E'; OURO = 'D2A93C'; CREME = 'F6F1E6'; CINZA = 'E1DCCF'
wb = Workbook(); ws = wb.active; ws.title = 'Variedades'

ws['A1'] = 'Aves Ornamentais Brasil — variedades por criadouro'
ws['A1'].font = Font(name='Arial', size=14, bold=True, color=VERDE)
ws['A2'] = ('Preencha a coluna E (Criadouro) com Aves Arca, Stima Aves ou Criadouro Aliança. '
            'As linhas em creme já estão definidas no site; as em amarelo são as que precisam ser separadas.')
ws['A2'].font = Font(name='Arial', size=9, italic=True, color='5B6B5B')
ws['A3'] = ('Quantidades e preços não entram aqui de propósito: variam todo dia. '
            'Esta planilha define só a QUEM pertence cada variedade.')
ws['A3'].font = Font(name='Arial', size=9, italic=True, color='5B6B5B')

HDR = ['#', 'Categoria', 'Grupo', 'Variedade', 'Criadouro', 'Observação (opcional)']
LIN = 5
for i, h in enumerate(HDR, start=1):
    c = ws.cell(row=LIN, column=i, value=h)
    c.font = Font(name='Arial', size=10, bold=True, color=CREME)
    c.fill = PatternFill('solid', fgColor=VERDE)
    c.alignment = Alignment(horizontal='left', vertical='center')
    c.border = Border(bottom=Side('thin', color=VERDE))
ws.row_dimensions[LIN].height = 22

fill_ok = PatternFill('solid', fgColor=CREME)
fill_edit = PatternFill('solid', fgColor='FFF6D6')
thin = Side('hair', color=CINZA)
NOME = {'aves-arca': 'Aves Arca'}
r = LIN + 1
for n, row in enumerate(rows, start=1):
    definido = row['criador'] == 'aves-arca'
    vals = [n, CAT.get(row['categoria'], row['categoria'].replace('-',' ').capitalize()), row['grupo'], row['nome'],
            NOME.get(row['criador'], ''), '']
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name='Arial', size=10, color='1E2A24',
                      bold=(i == 5 and definido))
        c.border = Border(bottom=thin)
        c.alignment = Alignment(vertical='center')
        if i in (5, 6):
            c.fill = fill_ok if definido else fill_edit
    r += 1
ULT = r - 1

dv = DataValidation(type='list', formula1='"Aves Arca,Stima Aves,Criadouro Aliança"',
                    allow_blank=True, showDropDown=False)
dv.error = 'Escolha um dos três criadouros da lista.'
dv.errorTitle = 'Criadouro inválido'
ws.add_data_validation(dv)
dv.add(f'E{LIN+1}:E{ULT}')

for col, w in zip('ABCDEF', (5, 13, 30, 34, 20, 34)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f'A{LIN+1}'
ws.auto_filter.ref = f'A{LIN}:F{ULT}'

# ---- Resumo com fórmulas (recalcula sozinho conforme você preenche) ----
rs = wb.create_sheet('Resumo')
rs['A1'] = 'Quantas variedades por criadouro'
rs['A1'].font = Font(name='Arial', size=12, bold=True, color=VERDE)
rs['A3'] = 'Criadouro'; rs['B3'] = 'Variedades'
for c in ('A3', 'B3'):
    rs[c].font = Font(name='Arial', size=10, bold=True, color=CREME)
    rs[c].fill = PatternFill('solid', fgColor=VERDE)
for i, nome in enumerate(['Aves Arca', 'Stima Aves', 'Criadouro Aliança'], start=4):
    rs[f'A{i}'] = nome
    rs[f'B{i}'] = f'=COUNTIF(Variedades!$E${LIN+1}:$E${ULT},A{i})'
    rs[f'A{i}'].font = Font(name='Arial', size=10)
    rs[f'B{i}'].font = Font(name='Arial', size=10)
rs['A7'] = 'Ainda sem criadouro'
rs['A7'].font = Font(name='Arial', size=10, bold=True, color='B5532E')
rs['B7'] = f'=COUNTBLANK(Variedades!$E${LIN+1}:$E${ULT})'
rs['B7'].font = Font(name='Arial', size=10, bold=True, color='B5532E')
rs['A9'] = 'Total de variedades'
rs['B9'] = f'=COUNTA(Variedades!$D${LIN+1}:$D${ULT})'
for c in ('A9', 'B9'):
    rs[c].font = Font(name='Arial', size=10, bold=True)
rs['A11'] = 'Por grupo — preencha o grupo em A12 e veja a divisão'
rs['A11'].font = Font(name='Arial', size=9, italic=True, color='5B6B5B')
rs['A12'] = 'Pavões'
rs['A12'].font = Font(name='Arial', size=10)
rs['A12'].fill = PatternFill('solid', fgColor='FFF6D6')
for i, nome in enumerate(['Aves Arca', 'Stima Aves', 'Criadouro Aliança'], start=13):
    rs[f'A{i}'] = nome
    rs[f'B{i}'] = (f'=COUNTIFS(Variedades!$C${LIN+1}:$C${ULT},$A$12,'
                   f'Variedades!$E${LIN+1}:$E${ULT},A{i})')
    rs[f'A{i}'].font = Font(name='Arial', size=10)
    rs[f'B{i}'].font = Font(name='Arial', size=10)
rs.column_dimensions['A'].width = 26
rs.column_dimensions['B'].width = 14

OUT = '/mnt/user-data/outputs/AOB - Variedades por criadouro - 07-09-2026.xlsx'
wb.save(OUT)
print('linhas', ULT - LIN, '->', OUT)
